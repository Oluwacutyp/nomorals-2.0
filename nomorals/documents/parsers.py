"""Parsers: bytes/path in any supported format -> :class:`Document`.

Format detection order: file extension first, then magic bytes
(PDF ``%PDF``, RTF ``{\\rtf``, ZIP ``PK`` for the OOXML/ODF/EPUB family
with content sniffing, OLE ``D0 CF 11 E0`` for legacy Office), then the
explicit ``mime`` hint, then a plain-text fallback.  Anything else raises
:class:`DocumentError` — the engine never returns a silently-empty
Document.

Use :func:`detect_format` to sniff a format without parsing.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from ..core.events import Event, global_bus
from ..core.logging_setup import get_logger
from ..core.pdf import PdfError, read_pdf_text
from ..tools.browser import node_to_markdown, parse_html
from .errors import DocumentError
from .model import Document, Section, Table, new_document
from .pdf_tables import extract_text_tables

__all__ = ["detect_format", "parse_bytes", "parse_path"]

_log = get_logger(__name__)


def _emit(topic: str, data: dict[str, Any]) -> None:
    """Publish a telemetry event. Best-effort: a broken bus or subscriber
    must never break parsing (fail-open telemetry, fail-closed function)."""
    try:
        global_bus.publish(Event(topic=topic, data=data, source=__name__))
    except Exception:  # noqa: BLE001 - telemetry is fail-open
        _log.debug("event %s failed", topic, exc_info=True)

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy Office compound document (.xls/.doc/.ppt)

_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xls": "xls",
    ".pptx": "pptx",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
    ".csv": "csv",
    ".tsv": "tsv",
    ".txt": "txt",
    ".rtf": "rtf",
    ".epub": "epub",
    ".odt": "odt",
    ".ods": "ods",
    ".json": "json",
    ".xml": "xml",
}

_MIMES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/rtf": "rtf",
    "text/rtf": "rtf",
    "application/epub+zip": "epub",
    "application/vnd.oasis.opendocument.text": "odt",
    "application/vnd.oasis.opendocument.spreadsheet": "ods",
    "application/json": "json",
    "text/json": "json",
    "application/xml": "xml",
    "text/xml": "xml",
    "text/html": "html",
    "text/markdown": "markdown",
    "text/x-markdown": "markdown",
    "text/csv": "csv",
    "text/tab-separated-values": "tsv",
    "text/plain": "txt",
}

# OOXML content-type markers used to tell docx/xlsx/pptx apart inside a ZIP.
_OOXML_MARKERS = (
    ("wordprocessingml", "docx"),
    ("spreadsheetml", "xlsx"),
    ("presentationml", "pptx"),
)

# Fallback part probing when [Content_Types].xml is missing or unhelpful.
_OOXML_PARTS = (
    ("word/document.xml", "docx"),
    ("xl/workbook.xml", "xlsx"),
    ("ppt/presentation.xml", "pptx"),
)

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
_HTML_START_RE = re.compile(r"(?is)^\s*(<!doctype\s+html|<html[\s>])")


def _stem(filename: str) -> str:
    name = Path(filename).name
    return name.rsplit(".", 1)[0] if "." in name else name


def _promote_title(doc: Document) -> None:
    """A leading h1 beats a filename-stem placeholder title."""
    if not doc.sections:
        return
    first = doc.sections[0]
    if first.level == 1 and first.heading:
        stem = _stem(doc.source)
        if not doc.title or doc.title == stem:
            doc.title = first.heading


def _sniff_zip(data: bytes) -> str:
    """Identify the ZIP-based flavour inside an archive."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            # OpenDocument (ODT/ODS): the spec mandates a leading "mimetype" entry.
            if "mimetype" in names:
                try:
                    mime = archive.read("mimetype").decode("ascii", "replace").strip()
                except Exception:  # noqa: BLE001 - fall through to other probes
                    mime = ""
                if mime == "application/vnd.oasis.opendocument.text":
                    return "odt"
                if mime == "application/vnd.oasis.opendocument.spreadsheet":
                    return "ods"
                if mime == "application/epub+zip":
                    return "epub"
            # EPUB: META-INF/container.xml points at the OPF package file.
            if "META-INF/container.xml" in names:
                try:
                    container = archive.read("META-INF/container.xml").decode(
                        "utf-8", "replace")
                    if ".opf" in container:
                        return "epub"
                except Exception:  # noqa: BLE001 - fall through to other probes
                    pass
            try:
                types_xml = archive.read("[Content_Types].xml").decode("utf-8", "replace")
            except KeyError:
                types_xml = ""
            for marker, fmt in _OOXML_MARKERS:
                if marker in types_xml:
                    return fmt
            for part, fmt in _OOXML_PARTS:
                if part in names:
                    return fmt
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"not a readable ZIP archive: {exc}") from exc
    raise DocumentError(
        "ZIP archive is not a recognized document "
        "(docx/xlsx/pptx/odt/ods/epub)")


def _sniff(data: bytes, filename: str, mime: str) -> str:
    """Return the format key for ``data`` or raise DocumentError."""
    if not data:
        raise DocumentError("cannot parse empty input")
    ext = Path(filename).suffix.lower() if filename else ""
    if ext in _EXTENSIONS:
        return _EXTENSIONS[ext]
    if ext:
        raise DocumentError(f"unsupported file extension {ext!r} for {filename!r}")
    if data[:4] == _PDF_MAGIC:
        return "pdf"
    if data[:4] == _ZIP_MAGIC:
        return _sniff_zip(data)
    if data[:4] == _OLE_MAGIC:
        # Legacy OLE compound document — could be .xls, .doc, or .ppt.
        # The container alone cannot tell them apart, so route by explicit
        # extension only and fail fast with a useful message otherwise.
        raise DocumentError(
            "legacy OLE Office document (could be .xls, .doc, or .ppt): "
            "pass filename= with the real extension so the right parser "
            "is chosen")
    if data.lstrip()[:5] == b"{\\rtf":
        return "rtf"
    if mime:
        fmt = _MIMES.get(mime.split(";")[0].strip().lower())
        if fmt:
            return fmt
        raise DocumentError(f"unsupported MIME type {mime!r}")
    text_start = data.lstrip()[:4096]
    if _HTML_START_RE.match(text_start.decode("utf-8", "replace")):
        return "html"
    # Content sniffing for extension-less JSON/XML (markitdown parity):
    # a bare '{"a": 1}' or '<root/>' with no filename should not become txt.
    stripped = text_start.decode("utf-8", "replace").strip()
    if stripped[:1] in ("{", "["):
        try:
            json.loads(stripped)
            return "json"
        except Exception:  # noqa: BLE001 - not JSON, keep sniffing
            pass
    if stripped[:1] == "<" and not stripped.startswith("</"):
        try:
            ET.fromstring(stripped)
            return "xml"
        except ET.ParseError:
            pass
    try:
        decoded = data.decode("utf-8")
    except UnicodeDecodeError:
        raise DocumentError(
            "unrecognized binary input: not PDF, not an Office document, "
            "not decodable text") from None
    if "\x00" in decoded:
        raise DocumentError("unrecognized binary input: NUL bytes in text stream")
    return "txt"


def detect_format(data: bytes, *, filename: str = "",
                  mime: str = "") -> str:
    """Return the format key for ``data`` without parsing it.

    Same detection as :func:`parse_bytes` (extension → magic bytes →
    MIME → text fallback).  Raises :class:`DocumentError` on empty or
    unrecognized input.
    """
    return _sniff(data, filename, mime)


# ── shared markdown sectioning ──────────────────────────────────────────────


def _parse_markdown_sections(text: str) -> list[Section]:
    """Split markdown-ish text into sections on ``#``..``######`` headings.

    Fenced code blocks pass through as body text — a ``#`` inside a fence
    is not a heading.
    """
    sections: list[Section] = []
    current = Section(level=1, heading="", text="")
    body: list[str] = []
    in_fence = False

    def flush() -> None:
        text_out = "\n".join(body).strip("\n")
        if current.heading or text_out.strip():
            sections.append(Section(level=current.level, heading=current.heading,
                                    text=text_out.strip()))
        body.clear()

    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_fence = not in_fence
            body.append(line)
            continue
        match = _HEADING_RE.match(stripped) if not in_fence else None
        if match:
            flush()
            current = Section(level=len(match.group(1)),
                              heading=match.group(2), text="")
        else:
            body.append(line)
    flush()
    if not sections:
        raise DocumentError("markdown document has no content")
    return sections


_LIST_ITEM_RE = re.compile(r"^(?:[-*•]|\d+[.)])\s+\S")


def _refine_kinds(sections: list[Section]) -> list[Section]:
    """Tag sections whose body is uniformly lists, quotes, or code.

    Parsers flatten structure to prose; this recovers the block kind so
    renderers and chunkers can treat it properly (docling keeps
    ListItem/CodeItem distinct for the same reason).
    """
    for section in sections:
        lines = [ln for ln in section.text.split("\n") if ln.strip()]
        if not lines or section.kind != "text":
            continue
        if all(_LIST_ITEM_RE.match(ln.strip()) for ln in lines):
            section.kind = "list"
        elif all(ln.strip().startswith(">") for ln in lines):
            section.kind = "quote"
        elif (len(lines) >= 1 and lines[0].strip().startswith("```")
              and lines[-1].strip().startswith("```")):
            section.kind = "code"
    return sections


# ── per-format parsers ──────────────────────────────────────────────────────


_PDF_INFO_KEY_RE = re.compile(rb"/(Title|Author|Subject|Keywords|Creator|Producer|CreationDate|ModDate)")
_PDF_LITERAL_RE = re.compile(rb"\((?:\\.|[^()\\])*\)")
_PDF_HEX_RE = re.compile(rb"<([0-9A-Fa-f\s]+)>")


def _pdf_unescape_string(raw: bytes) -> str:
    """Decode a PDF literal string ``( ... )`` (escapes + octal codes)."""
    body = raw[1:-1]
    out: list[str] = []
    i, n = 0, len(body)
    simple = {ord("n"): "\n", ord("r"): "\r", ord("t"): "\t", ord("b"): "\b",
              ord("f"): "\f", ord("\\"): "\\", ord("("): "(", ord(")"): ")"}
    while i < n:
        byte = body[i]
        if byte != 0x5C or i + 1 >= n:  # not a backslash escape
            out.append(chr(byte))
            i += 1
            continue
        nxt = body[i + 1]
        if nxt in simple:
            out.append(simple[nxt])
            i += 2
        elif nxt in (0x0A, 0x0D):  # line continuation: drop
            i += 2
            if nxt == 0x0D and i < n and body[i] == 0x0A:
                i += 1
        elif 0x30 <= nxt <= 0x37:  # octal \ddd
            digits = body[i + 1:i + 4]
            count = 0
            for digit in digits:
                if 0x30 <= digit <= 0x37:
                    count += 1
                else:
                    break
            out.append(chr(int(digits[:count], 8) & 0xFF))
            i += 1 + count
        else:
            out.append(chr(nxt))
            i += 2
    return "".join(out)


def _pdf_info(data: bytes) -> dict[str, str]:
    """Extract the PDF /Info dictionary (title/author/dates), best effort.

    Returns {} when the trailer has no /Info entry or it cannot be read —
    metadata is a bonus, never a failure.
    """
    info_ref = re.search(rb"/Info\s+(\d+)\s+\d+\s+R", data)
    if not info_ref:
        return {}
    obj = re.search(rb"\b" + info_ref.group(1) + rb"\s+\d+\s+obj(.*?)endobj",
                    data, re.DOTALL)
    if not obj:
        return {}
    body = obj.group(1)
    out: dict[str, str] = {}
    for key_match in _PDF_INFO_KEY_RE.finditer(body):
        key = key_match.group(1).decode("ascii")
        rest = body[key_match.end():]
        literal = _PDF_LITERAL_RE.match(rest.lstrip())
        if literal is not None:
            out[key] = _pdf_unescape_string(literal.group(0))
            continue
        hexed = _PDF_HEX_RE.match(rest.lstrip())
        if hexed is not None:
            try:
                raw_hex = re.sub(rb"\s", b"", hexed.group(1))
                decoded = bytes.fromhex(raw_hex.decode("ascii"))
                if decoded[:2] == b"\xfe\xff":
                    out[key] = decoded[2:].decode("utf-16-be", "replace")
                else:
                    out[key] = decoded.decode("latin-1", "replace")
            except (ValueError, UnicodeDecodeError):
                # A corrupt hex Info value is best-effort metadata, not a
                # parse failure: skip the key and keep the document.
                _log.debug("unreadable hex PDF Info value for %s", key)
    for date_key in ("CreationDate", "ModDate"):
        raw = out.get(date_key, "")
        parsed = _pdf_date_to_iso(raw)
        if parsed:
            out[date_key] = parsed
    return out


def _pdf_date_to_iso(raw: str) -> str:
    """``D:20261003120000+02'00'`` -> ``2026-10-03T12:00:00``; '' on failure."""
    match = re.match(
        r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", raw.strip())
    if not match:
        return ""
    year, mon, day, hour, minute, second = match.groups()
    try:
        return (f"{year}-{mon or '01'}-{day or '01'}"
                f"T{hour or '00'}:{minute or '00'}:{second or '00'}")
    except (TypeError, ValueError):
        return ""


def _parse_pdf(data: bytes, doc: Document) -> Document:
    try:
        text = read_pdf_text(data)
    except PdfError as exc:
        raise DocumentError(f"invalid PDF: {exc}") from exc
    if not text.strip():
        raise DocumentError("PDF contains no extractable text (scanned/image-only?)")
    # read_pdf_text joins pages with "\n\n" and never emits a blank line
    # inside a page, so splitting there recovers the page boundaries.
    pages = [p for p in text.split("\n\n") if p.strip()]
    doc.sections = [Section(level=1, heading=f"Page {i + 1}", text=page,
                            page=i + 1)
                    for i, page in enumerate(pages)]
    # Recover whitespace-aligned tables from the page text (conservative:
    # only consistent columnar blocks become tables).
    for i, page in enumerate(pages):
        for table in extract_text_tables(page):
            table.name = f"Page {i + 1} {table.name}"
            table.page = i + 1
            doc.tables.append(table)
    doc.metadata["pages"] = len(pages)
    if doc.tables:
        doc.metadata["tables"] = len(doc.tables)
    # PDF document-information dictionary: title/author/dates.
    info = _pdf_info(data)
    if info.get("Title"):
        doc.title = info["Title"]
    if info.get("Author"):
        doc.author = info["Author"]
    for key in ("Subject", "Keywords", "Creator", "Producer",
                "CreationDate", "ModDate"):
        if info.get(key):
            doc.metadata[key.lower()] = info[key]
    return doc


_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_NUM_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_hyperlink_targets(data: bytes) -> dict[str, str]:
    """Map ``rId`` → URL for ``w:hyperlink`` elements (best effort)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            raw = archive.read("word/_rels/document.xml.rels")
        root = ET.fromstring(raw)
    except Exception:  # noqa: BLE001 - links are a bonus, not a failure
        return {}
    targets: dict[str, str] = {}
    for rel in root.iter(f"{_PKG_REL_NS}Relationship"):
        rel_type = rel.get("Type", "")
        if rel_type.endswith("/hyperlink"):
            rid, target = rel.get("Id", ""), rel.get("Target", "")
            if rid and target:
                targets[rid] = target
    return targets


def _docx_list_kinds(data: bytes) -> dict[str, str]:
    """Map ``w:numId`` → ``"bullet"`` or ``"decimal"`` via numbering.xml."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            raw = archive.read("word/numbering.xml")
        root = ET.fromstring(raw)
    except Exception:  # noqa: BLE001 - lists degrade to bullets
        return {}
    abstract_fmt: dict[str, str] = {}
    for abstract in root.iter(f"{_NUM_NS}abstractNum"):
        aid = abstract.get(f"{_NUM_NS}abstractNumId", "")
        fmt = "bullet"
        for lvl in abstract.iter(f"{_NUM_NS}lvl"):
            num_fmt = lvl.find(f"{_NUM_NS}numFmt")
            if num_fmt is not None:
                val = num_fmt.get(f"{_NUM_NS}val", "")
                fmt = "decimal" if val == "decimal" else "bullet"
                break
        if aid:
            abstract_fmt[aid] = fmt
    kinds: dict[str, str] = {}
    for num in root.iter(f"{_NUM_NS}num"):
        num_id = num.get(f"{_NUM_NS}numId", "")
        ref = num.find(f"{_NUM_NS}abstractNumId")
        aid = ref.get(f"{_NUM_NS}val", "") if ref is not None else ""
        if num_id:
            kinds[num_id] = abstract_fmt.get(aid, "bullet")
    return kinds


def _docx_footnotes(data: bytes) -> list[tuple[str, str]]:
    """(id, text) footnotes from ``word/footnotes.xml`` (best effort)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            raw = archive.read("word/footnotes.xml")
        root = ET.fromstring(raw)
    except Exception:  # noqa: BLE001 - footnotes are a bonus
        return []
    notes: list[tuple[str, str]] = []
    for note in root.iter(f"{_NUM_NS}footnote"):
        if note.get(f"{_NUM_NS}type", "") in (
                "separator", "continuationSeparator", "continuationNotice"):
            continue
        fid = note.get(f"{_NUM_NS}id", "")
        text = " ".join(
            "".join(t.text or "" for t in p.iter(f"{_NUM_NS}t")).strip()
            for p in note.iter(f"{_NUM_NS}p")).strip()
        if fid and text:
            notes.append((fid, text))
    return notes


def _para_num_id(child: Any, para: Any, word_ns: str) -> str:
    """Numbering id for a docx paragraph: direct ``w:numPr`` first, then
    the style chain (python-docx's List styles number via the style
    definition, real Word files usually number directly)."""
    p_pr = child.find(f"{word_ns}pPr")
    if p_pr is not None:
        num_pr = p_pr.find(f"{word_ns}numPr")
        if num_pr is not None:
            num_id = num_pr.find(f"{word_ns}numId")
            if num_id is not None:
                return num_id.get(f"{word_ns}val", "")
    seen: set[int] = set()
    style = getattr(para, "style", None)
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        element = getattr(style, "element", None)
        if element is not None:
            for num_pr in element.iter(f"{word_ns}numPr"):
                num_id = num_pr.find(f"{word_ns}numId")
                if num_id is not None:
                    return num_id.get(f"{word_ns}val", "")
        try:
            style = style.base_style
        except Exception:  # noqa: BLE001 - end of the style chain
            break
    return ""


def _parse_docx(data: bytes, doc: Document) -> Document:
    try:
        from docx import Document as DocxDocument
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentError(
            "parsing .docx requires the optional 'python-docx' package "
            "(pip install python-docx)") from exc

    try:
        oxml_doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"invalid .docx file: {exc}") from exc
    props = oxml_doc.core_properties
    if props.title:
        doc.title = props.title
    if props.author:
        doc.author = props.author
    for attr, key in (("keywords", "keywords"), ("comments", "comments"),
                      ("description", "description"),
                      ("last_modified_by", "last_modified_by")):
        value = getattr(props, attr, None)
        if value:
            doc.metadata[key] = str(value)
    for attr, key in (("created", "created"), ("modified", "modified")):
        value = getattr(props, attr, None)
        if value is not None and hasattr(value, "isoformat"):
            doc.metadata[key] = value.isoformat()

    word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    sections: list[Section] = []
    current = Section(level=1, heading="", text="")
    body: list[str] = []
    table_count = 0
    link_targets = _docx_hyperlink_targets(data)
    list_kinds = _docx_list_kinds(data)
    links: list[str] = []
    list_counters: dict[str, int] = {}

    def flush_section() -> None:
        text_out = "\n".join(body).strip()
        if current.heading or text_out:
            sections.append(Section(level=current.level, heading=current.heading,
                                    text=text_out))
        body.clear()

    for child in oxml_doc.element.body.iterchildren():
        if child.tag == f"{word_ns}p":
            para = Paragraph(child, oxml_doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            if style == "Title":
                # A Title style beats the filename-stem placeholder.
                stem = _stem(doc.source)
                if not doc.title or doc.title == stem:
                    doc.title = text
                continue
            if style == "Subtitle":
                doc.metadata.setdefault("subtitle", text)
                continue
            # Hyperlinks: keep the visible text (already in para.text)
            # and harvest the URLs into metadata.
            for hyperlink in child.iter(f"{word_ns}hyperlink"):
                rid = hyperlink.get(f"{_REL_NS}id", "")
                target = link_targets.get(rid, "")
                if target and target not in links:
                    links.append(target)
            if style.startswith("Heading"):
                flush_section()
                try:
                    level = int(style.rsplit(" ", 1)[1])
                except (ValueError, IndexError):
                    level = 1
                current = Section(level=min(max(level, 1), 6), heading=text, text="")
                continue
            # Lists: w:numPr (direct or via the style chain) or a List
            # style → bullet/numbered marker.  markitdown preserves list
            # structure; flat paragraphs lose it.
            num_id = _para_num_id(child, para, word_ns)
            is_list = bool(num_id) or style.startswith("List")
            if is_list:
                kind = list_kinds.get(num_id, "bullet")
                if kind == "decimal":
                    list_counters[num_id] = list_counters.get(num_id, 0) + 1
                    text = f"{list_counters[num_id]}. {text}"
                else:
                    text = f"• {text}"
            body.append(text)
        elif child.tag == f"{word_ns}tbl":
            table_count += 1
            dtable = DocxTable(child, oxml_doc)
            cells = [[c.text.strip() for c in row.cells] for row in dtable.rows]
            cells = [r for r in cells if any(r)]
            if not cells:
                continue
            name = f"Table {table_count}"
            doc.tables.append(Table(name=name, headers=cells[0], rows=cells[1:]))
            body.append(f"[{name}: {len(cells) - 1} data rows]")
    flush_section()
    if links:
        doc.metadata["links"] = links
    for fid, ftext in _docx_footnotes(data):
        sections.append(Section(level=2, heading=f"Footnote [{fid}]",
                                text=ftext, kind="footnote"))
    if not sections and not doc.tables:
        raise DocumentError(".docx contains no paragraphs or tables")
    doc.sections = _refine_kinds(sections)
    doc.metadata["tables"] = len(doc.tables)
    return doc


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()  # type: ignore[union-attr]
        except Exception:  # noqa: BLE001 - fall back to str() for odd date-likes
            pass
    return str(value).strip()


def _parse_xlsx(data: bytes, doc: Document) -> Document:
    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentError(
            "parsing .xlsx requires the optional 'openpyxl' package "
            "(pip install openpyxl)") from exc

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                          data_only=True)
    except Exception as exc:
        raise DocumentError(f"invalid .xlsx file: {exc}") from exc
    props = workbook.properties
    if getattr(props, "title", None):
        doc.title = props.title
    if getattr(props, "creator", None):
        doc.author = props.creator
    for attr, key in (("keywords", "keywords"), ("description", "description"),
                      ("lastModifiedBy", "last_modified_by")):
        value = getattr(props, attr, None)
        if value:
            doc.metadata[key] = str(value)
    for attr, key in (("created", "created"), ("modified", "modified")):
        value = getattr(props, attr, None)
        if value is not None and hasattr(value, "isoformat"):
            doc.metadata[key] = value.isoformat()

    for sheet in workbook.worksheets:
        grid = [[_cell_text(cell.value) for cell in row]
                for row in sheet.iter_rows()]
        grid = [row for row in grid if any(cell for cell in row)]
        if not grid:
            continue
        width = max(len(row) for row in grid)
        padded = [row + [""] * (width - len(row)) for row in grid]
        headers, rows = padded[0], padded[1:]
        name = sheet.title or f"Sheet {len(doc.tables) + 1}"
        doc.tables.append(Table(name=name, headers=headers, rows=rows))
        doc.sections.append(Section(
            level=1, heading=name,
            text=(f"Worksheet {name!r}: {len(rows)} data rows × "
                  f"{len(headers)} columns.")))
        merged = getattr(sheet, "merged_cells", None)
        merged_count = len(list(merged.ranges)) if merged is not None else 0
        if merged_count:
            doc.metadata.setdefault("merged_cells", 0)
            doc.metadata["merged_cells"] += merged_count
    workbook.close()
    if not doc.tables:
        raise DocumentError(".xlsx workbook has no non-empty worksheets")
    doc.metadata["sheets"] = [t.name for t in doc.tables]
    return doc


_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_A = f"{{{_DRAWING_NS}}}"


def _parse_pptx(data: bytes, doc: Document) -> Document:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"invalid .pptx file: {exc}") from exc
    with archive:
        numbered: list[tuple[int, str]] = []
        for entry in archive.namelist():
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", entry)
            if match:
                numbered.append((int(match.group(1)), entry))
        if not numbered:
            raise DocumentError(".pptx contains no slides")
        numbered.sort()

        for index, (_num, name) in enumerate(numbered, 1):
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError as exc:
                raise DocumentError(f"malformed slide XML in {name}: {exc}") from exc
            # a:t runs that live inside a:tbl cells are rendered as table
            # rows instead, so collect their element ids to skip them below.
            table_text_ids: set[int] = set()
            table_lines: list[str] = []
            for tbl in root.iter(f"{_A}tbl"):
                for t_elem in tbl.iter(f"{_A}t"):
                    table_text_ids.add(id(t_elem))
                for tr in tbl.findall(f"{_A}tr"):
                    cells = []
                    for tc in tr.findall(f"{_A}tc"):
                        cell = " ".join(t.text or "" for t in tc.iter(f"{_A}t"))
                        cells.append(cell.strip())
                    if any(cells):
                        table_lines.append(" | ".join(cells))
            lines = [(elem.text or "").strip()
                     for elem in root.iter(f"{_A}t")
                     if id(elem) not in table_text_ids]
            lines = [line for line in lines if line]
            lines.extend(table_lines)
            if not lines:
                continue
            doc.sections.append(Section(level=1, heading=f"Slide {index}",
                                        text="\n".join(lines), page=index))
    if not doc.sections:
        raise DocumentError(".pptx slides contain no text")
    doc.metadata["slides"] = len(doc.sections)
    return doc


def _parse_html(data: bytes, doc: Document) -> Document:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentError(f"HTML is not valid UTF-8: {exc}") from exc
    root = parse_html(text)
    titles = root.find_all("title")
    if titles and titles[0].inner_text().strip():
        doc.title = titles[0].inner_text().strip()
    # <head> metadata harvest: description/keywords/author + lang.
    for meta in root.find_all("meta"):
        name = (meta.attrs.get("name", "") or "").lower()
        content = meta.attrs.get("content", "").strip()
        if not content:
            continue
        if name in ("description", "keywords", "author", "creator",
                    "subject"):
            doc.metadata[name] = content
    for html_node in root.find_all("html"):
        lang = html_node.attrs.get("lang", "").strip()
        if lang:
            doc.metadata["language"] = lang
            break
    # <head> is metadata: drop it so its text can't glue itself to the
    # first body heading (node_to_markdown strips inner block edges).
    for head in root.find_all("head"):
        parent = head.parent
        if parent is not None and head in parent.children:
            parent.children.remove(head)
    bodies = root.find_all("body")
    markdown = node_to_markdown(bodies[0] if bodies else root)
    if not markdown.strip():
        raise DocumentError("HTML document has no readable content")
    doc.sections = _refine_kinds(_parse_markdown_sections(markdown))
    _promote_title(doc)
    for i, table_node in enumerate(root.find_all("table"), 1):
        grid: list[list[str]] = []
        for tr in table_node.find_all("tr"):
            cells = [c.inner_text().strip()
                     for c in tr.find_all(("th", "td"))]
            if any(cells):
                grid.append(cells)
        if not grid:
            continue
        width = max(len(row) for row in grid)
        grid = [row + [""] * (width - len(row)) for row in grid]
        doc.tables.append(Table(name=f"Table {i}", headers=grid[0], rows=grid[1:]))
    return doc


def _parse_markdown(data: bytes, doc: Document) -> Document:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentError(f"markdown is not valid UTF-8: {exc}") from exc
    if not text.strip():
        raise DocumentError("markdown document is empty")
    doc.sections = _refine_kinds(_parse_markdown_sections(text))
    _promote_title(doc)
    return doc


def _parse_delimited(data: bytes, doc: Document, delimiter: str, label: str) -> Document:
    # Encoding fallback chain: UTF-8 (with BOM) → Windows cp1252 →
    # latin-1 (never fails).  Non-UTF-8 CSVs are common in the wild;
    # dying on them is worse than a best-effort decode.
    text = ""
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise DocumentError(f"{label} file is empty")
    if delimiter == ",":
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid = [[cell.strip() for cell in row] for row in reader]
    grid = [row for row in grid if any(cell for cell in row)]
    if not grid:
        raise DocumentError(f"{label} file has no data rows")
    width = max(len(row) for row in grid)
    grid = [row + [""] * (width - len(row)) for row in grid]
    name = _stem(doc.source) or "data"
    doc.tables.append(Table(name=name, headers=grid[0], rows=grid[1:]))
    doc.sections.append(Section(
        level=1, heading=name,
        text=f"{label.upper()} table {name!r}: {len(grid) - 1} data rows × "
             f"{len(grid[0])} columns."))
    doc.metadata["delimiter"] = delimiter
    return doc


def _parse_txt(data: bytes, doc: Document) -> Document:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentError(f"text is not valid UTF-8: {exc}") from exc
    if not text.strip():
        raise DocumentError("text document is empty")
    doc.sections = [Section(level=1, heading="", text=text.strip())]
    return doc


# ── JSON / XML (stdlib) ───────────────────────────────────────────────────

_JSON_MAX_DEPTH = 6


def _json_walk(value: Any, heading: str, level: int, doc: Document,
               depth: int) -> None:
    """Recursively lower a JSON value into sections/tables."""
    if depth > _JSON_MAX_DEPTH:
        doc.sections.append(Section(level=min(level, 6), heading=heading,
                                    text="[nested too deep — truncated]"))
        return
    level = min(max(level, 1), 6)
    if isinstance(value, dict):
        if not value:
            return
        # A dict of scalars becomes prose lines; anything richer recurses.
        scalars = {k: v for k, v in value.items()
                   if isinstance(v, (str, int, float, bool)) or v is None}
        rest = {k: v for k, v in value.items() if k not in scalars}
        body = "\n".join(f"{k}: {v}" for k, v in scalars.items()
                         if v is not None and str(v).strip())
        if body or not rest:
            doc.sections.append(Section(level=level, heading=heading,
                                        text=body))
        for key, sub in rest.items():
            _json_walk(sub, str(key), level + 1, doc, depth + 1)
    elif isinstance(value, list):
        if not value:
            return
        if all(isinstance(item, dict) for item in value):
            keys: list[str] = []
            for item in value:
                for key in item:
                    if key not in keys:
                        keys.append(str(key))
            rows = [[str(item.get(k, "")) for k in keys] for item in value]
            doc.tables.append(Table(name=heading or f"Table {len(doc.tables) + 1}",
                                    headers=keys, rows=rows))
            doc.sections.append(Section(
                level=level, heading=heading,
                text=f"JSON array {heading!r}: {len(rows)} records × "
                     f"{len(keys)} fields."))
        elif all(isinstance(item, (str, int, float, bool)) or item is None
                 for item in value):
            items = [str(item) for item in value if item is not None]
            doc.sections.append(Section(
                level=level, heading=heading, kind="list",
                text="\n".join(f"• {item}" for item in items)))
        else:
            for i, item in enumerate(value):
                _json_walk(item, f"{heading} [{i + 1}]" if heading
                           else f"[{i + 1}]", level, doc, depth + 1)
    elif value is not None:
        text = str(value).strip()
        if text:
            doc.sections.append(Section(level=level, heading=heading,
                                        text=text))


def _parse_json(data: bytes, doc: Document) -> Document:
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DocumentError(f"invalid JSON: {exc}") from exc
    before = (len(doc.sections), len(doc.tables))
    if isinstance(payload, dict):
        for key, value in payload.items():
            _json_walk(value, str(key), 1, doc, 0)
    elif isinstance(payload, list):
        _json_walk(payload, _stem(doc.source) or "data", 1, doc, 0)
    else:
        _json_walk(payload, "", 1, doc, 0)
    if (len(doc.sections), len(doc.tables)) == before:
        raise DocumentError("JSON document has no readable content")
    if doc.tables:
        doc.metadata["tables"] = len(doc.tables)
    return doc


_XML_MAX_DEPTH = 6


def _xml_text(elem: Any) -> str:
    return "".join(elem.itertext()).strip()


def _xml_walk(elem: Any, level: int, doc: Document, depth: int) -> None:
    """One top-level XML element → section; nested elements → subsections."""
    if depth > _XML_MAX_DEPTH:
        return
    level = min(max(level, 1), 6)
    tag = _localname(elem.tag)
    children = [c for c in elem if isinstance(c.tag, str)]
    attr_lines = [f"@{k}: {v}" for k, v in elem.attrib.items()]
    if not children:
        body = "\n".join(attr_lines)
        text = _xml_text(elem)
        if text:
            body = (body + "\n" + text).strip() if body else text
        if body or attr_lines:
            doc.sections.append(Section(level=level, heading=tag, text=body))
        return
    # Mixed element: its own direct text first, then child subsections.
    direct = (elem.text or "").strip()
    tails = " ".join((c.tail or "").strip() for c in children).strip()
    own = " ".join(p for p in (direct, tails) if p)
    if attr_lines:
        own = ("\n".join(attr_lines) + ("\n" + own if own else "")).strip()
    if own:
        doc.sections.append(Section(level=level, heading=tag, text=own))
    for child in children:
        _xml_walk(child, level + 1, doc, depth + 1)


def _parse_xml(data: bytes, doc: Document) -> Document:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentError(f"XML is not valid UTF-8: {exc}") from exc
    if not text.strip():
        raise DocumentError("XML document is empty")
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise DocumentError(f"invalid XML: {exc}") from exc
    doc.metadata["xml_root"] = _localname(root.tag)
    if root.attrib:
        doc.metadata["xml_attributes"] = dict(root.attrib)
    for child in root:
        if isinstance(child.tag, str):
            _xml_walk(child, 1, doc, 0)
    # A root with only direct text (no element children).
    if not doc.sections:
        body = _xml_text(root)
        if body:
            doc.sections.append(Section(level=1, heading=_localname(root.tag),
                                        text=body))
    if not doc.sections:
        raise DocumentError("XML document has no readable content")
    return doc


def _xls_cell_text(cell: Any, datemode: int) -> str:
    """Render an xlrd cell to string (dates → ISO, like _cell_text)."""
    import datetime as _dt
    ctype = cell.ctype
    value = cell.value
    if ctype == 0 or value == "":  # empty
        return ""
    if ctype == 2:  # number
        if float(value).is_integer():
            return str(int(value))
        return str(value)
    if ctype == 3:  # date
        try:
            import xlrd
            dt = xlrd.xldate_as_datetime(value, datemode)
            if isinstance(dt, _dt.datetime):
                return dt.isoformat()
            return str(dt)
        except Exception:  # noqa: BLE001 - fall back to raw value
            return str(value)
    if ctype == 4:  # boolean
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _parse_xls(data: bytes, doc: Document) -> Document:
    try:
        import xlrd
    except ImportError as exc:
        raise DocumentError(
            "parsing .xls requires the optional 'xlrd' package "
            "(pip install xlrd)") from exc
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise DocumentError(f"invalid .xls file: {exc}") from exc
    for sheet in book.sheets():
        grid = [[_xls_cell_text(sheet.cell(r, c), book.datemode)
                 for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
        grid = [row for row in grid if any(cell for cell in row)]
        if not grid:
            continue
        width = max(len(row) for row in grid)
        padded = [row + [""] * (width - len(row)) for row in grid]
        headers, rows = padded[0], padded[1:]
        name = sheet.name or f"Sheet {len(doc.tables) + 1}"
        doc.tables.append(Table(name=name, headers=headers, rows=rows))
        doc.sections.append(Section(
            level=1, heading=name,
            text=(f"Worksheet {name!r}: {len(rows)} data rows × "
                  f"{len(headers)} columns.")))
    if not doc.tables:
        raise DocumentError(".xls workbook has no non-empty worksheets")
    doc.metadata["sheets"] = [t.name for t in doc.tables]
    return doc


# ── RTF (stdlib-only) ───────────────────────────────────────────────────────

#: Destination groups whose content is formatting, not prose.
_RTF_SKIP_DESTINATIONS = frozenset({
    "fonttbl", "colortbl", "stylesheet", "info", "themedata", "listtable",
    "listoverridetable", "rsidtbl", "generator", "colorschememapping",
    "datastore", "header", "footer", "footnote", "annotation", "pict",
    "shppict", "pnseclvl", "themedata", "latentstyles",
})

#: One-shot control words that map to literal characters.
_RTF_SPECIAL = {
    "par": "\n", "line": "\n", "tab": "\t", "emdash": "\u2014",
    "endash": "\u2013", "bullet": "\u2022", "lquote": "\u2018",
    "rquote": "\u2019", "ldblquote": "\u201c", "rdblquote": "\u201d",
}

_RTF_WORD_RE = re.compile(r"[a-z]+")
_RTF_NUM_RE = re.compile(r"-?\d+")


def _rtf_to_text(data: bytes) -> str:
    """Extract plain text from RTF markup.

    Handles groups (with ignorable ``{\\*...}`` destinations and known
    formatting destinations skipped), control words, ``\\uN`` Unicode
    escapes with their fallback character, and ``\\'hh`` hex escapes.
    """
    src = data.decode("latin-1", "replace")  # 1:1 bytes; \'hh handled below
    out: list[str] = []
    skip_stack: list[bool] = []  # skip-state of each open group
    skip = False
    i, n = 0, len(src)
    while i < n:
        ch = src[i]
        if ch == "{":
            skip_stack.append(skip)
            i += 1
            continue
        if ch == "}":
            skip = skip_stack.pop() if skip_stack else False
            i += 1
            continue
        if ch != "\\":
            # Raw newlines in RTF source are formatting, not content.
            if not skip and ch not in "\r\n":
                out.append(ch)
            i += 1
            continue
        i += 1
        if i >= n:
            break
        esc = src[i]
        if esc in "\\{}":
            if not skip:
                out.append(esc)
            i += 1
            continue
        if esc == "'":
            hexpair = src[i + 1:i + 3]
            try:
                code = int(hexpair, 16)
            except ValueError:
                code = 0x3F  # '?'
            if not skip:
                out.append(chr(code))
            i += 3
            continue
        if esc == "*":  # ignorable destination group
            skip = True
            i += 1
            continue
        if esc == "~":
            if not skip:
                out.append("\u00a0")
            i += 1
            continue
        if esc in "-_":
            # \- optional hyphen (dropped); \_ non-breaking hyphen
            if not skip and esc == "_":
                out.append("-")
            i += 1
            continue
        word_match = _RTF_WORD_RE.match(src, i)
        if not word_match:
            i += 1
            continue
        word = word_match.group(0)
        i = word_match.end()
        num: int | None = None
        num_match = _RTF_NUM_RE.match(src, i)
        if num_match:
            num = int(num_match.group(0))
            i = num_match.end()
        if i < n and src[i] == " ":  # control-word delimiter space
            i += 1
        if skip:
            continue
        if word in _RTF_SKIP_DESTINATIONS:
            skip = True
            continue
        if word == "u" and num is not None:
            out.append(chr(num + 65536 if num < 0 else num))
            # Skip the single fallback character (a char or \'hh).
            if i < n:
                if src[i] == "\\" and i + 3 < n and src[i + 1] == "'":
                    i += 4
                else:
                    i += 1
            continue
        special = _RTF_SPECIAL.get(word)
        if special is not None:
            out.append(special)
    return "".join(out)


def _parse_rtf(data: bytes, doc: Document) -> Document:
    src = data.decode("latin-1", "replace")
    tables, clean_src = _rtf_extract_tables(src)
    doc.tables.extend(tables)
    if tables:
        doc.metadata["tables"] = len(tables)
    text = _rtf_to_text(clean_src.encode("latin-1", "replace"))
    blocks = [b.strip() for b in re.split(r"\n{2,}|\r\n{2,}", text) if b.strip()]
    # Single-newline paragraphs collapse into blocks; stray lone lines that
    # look like headings are not promoted — RTF carries no heading info.
    if not blocks and not tables:
        raise DocumentError("RTF document contains no readable text")
    doc.sections = [Section(level=1, heading="", text=block)
                    for block in blocks]
    return doc


# ── RTF tables (``\trowd … \cell … \row``) ─────────────────────────────────

_RTF_ROW_RE = re.compile(r"\\trowd\b(.*?)\\row\b", re.DOTALL)
_RTF_HEX_RE = re.compile(r"\\'[0-9a-fA-F]{2}")
_RTF_UNI_RE = re.compile(r"\\u(-?\d+)")
_RTF_CTRL_RE = re.compile(r"\\[a-z]+\d* ?")


def _rtf_fragment_text(fragment: str) -> str:
    """Best-effort plain text of an RTF cell fragment (controls stripped)."""
    text = _RTF_HEX_RE.sub(
        lambda m: chr(int(m.group(0)[2:], 16)), fragment)
    text = _RTF_UNI_RE.sub(
        lambda m: chr(int(m.group(1)) % 65536), text)
    text = _RTF_CTRL_RE.sub(" ", text)
    text = text.replace("{", " ").replace("}", " ")
    text = text.replace("\\\\", "\\")
    return re.sub(r"\s+", " ", text).strip()


def _rtf_extract_tables(src: str) -> tuple[list[Table], str]:
    """Recover ``\\trowd`` tables from RTF source.

    Returns (tables, cleaned_source): consecutive ``\\trowd…\\row`` runs
    become Tables (first row = header), and their spans are blanked from
    the source so the prose pass doesn't duplicate them.
    """
    row_spans: list[tuple[int, int, list[str]]] = []
    for match in _RTF_ROW_RE.finditer(src):
        parts = re.split(r"\\cell\b", match.group(1))
        cells = [_rtf_fragment_text(part) for part in parts]
        while cells and not cells[-1]:
            cells.pop()  # trailing row formatting after the last \cell
        if cells and any(cells):
            row_spans.append((match.start(), match.end(), cells))
    grouped: list[list[list[str]]] = []
    current: list[list[str]] = []
    prev_end: int | None = None
    for start, end, cells in row_spans:
        if prev_end is not None and src[prev_end:start].strip():
            if current:
                grouped.append(current)
                current = []
        current.append(cells)
        prev_end = end
    if current:
        grouped.append(current)
    tables: list[Table] = []
    for rows in grouped:
        width = max(len(row) for row in rows)
        grid = [row + [""] * (width - len(row)) for row in rows]
        tables.append(Table(name=f"Table {len(tables) + 1}",
                            headers=grid[0], rows=grid[1:],
                            confidence=0.9))
    clean = src
    for start, end, _cells in reversed(row_spans):
        clean = clean[:start] + "\n" + clean[end:]
    return tables, clean


# ── EPUB (stdlib-only zip + XML) ────────────────────────────────────────────

_OPF_NS = "{http://www.idpf.org/2007/opf}"
_DC_NS = "{http://purl.org/dc/elements/1.1/}"
_CONTAINER_NS = "{urn:oasis:names:tc:opendocument:xmlns:container}"
_EPUB_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}


def _localname(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _xhtml_text(elem: Any) -> str:
    return "".join(elem.itertext()).strip()


def _parse_epub(data: bytes, doc: Document) -> Document:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"invalid .epub file: {exc}") from exc
    with archive:
        opf_path = ""
        try:
            container = ET.fromstring(archive.read("META-INF/container.xml"))
        except (KeyError, ET.ParseError):
            # Fallback: find .opf file directly in the archive
            for name in archive.namelist():
                if name.endswith(".opf"):
                    opf_path = name
                    break
        else:
            rootfile = container.find(f".//{_CONTAINER_NS}rootfile")
            opf_path = (rootfile.get("full-path", "")
                        if rootfile is not None else "")
        if not opf_path or opf_path not in archive.namelist():
            raise DocumentError(".epub has no OPF package file")
        try:
            opf = ET.fromstring(archive.read(opf_path))
        except ET.ParseError as exc:
            raise DocumentError(f"malformed OPF package file {opf_path}: {exc}") from exc
        title_node = opf.find(f".//{_OPF_NS}metadata/{_DC_NS}title")
        if title_node is not None and title_node.text and title_node.text.strip():
            doc.title = title_node.text.strip()
        creator_node = opf.find(f".//{_OPF_NS}metadata/{_DC_NS}creator")
        if (creator_node is not None and creator_node.text
                and creator_node.text.strip()):
            doc.author = creator_node.text.strip()
        lang_node = opf.find(f".//{_OPF_NS}metadata/{_DC_NS}language")
        if (lang_node is not None and lang_node.text
                and lang_node.text.strip()):
            doc.metadata["language"] = lang_node.text.strip()
        subjects = [n.text.strip() for n in
                    opf.findall(f".//{_OPF_NS}metadata/{_DC_NS}subject")
                    if n.text and n.text.strip()]
        if subjects:
            doc.metadata["subjects"] = subjects
        manifest = {item.get("id", ""): item.get("href", "")
                    for item in opf.iter(f"{_OPF_NS}item")}
        opf_dir = str(Path(opf_path).parent)
        spine_hrefs: list[str] = []
        for ref in opf.iter(f"{_OPF_NS}itemref"):
            href = manifest.get(ref.get("idref", ""), "")
            if not href:
                continue
            full = href if opf_dir in ("", ".") else f"{opf_dir}/{href}"
            if full in archive.namelist() and \
                    full.lower().endswith((".xhtml", ".html", ".htm")):
                spine_hrefs.append(full)

        sections: list[Section] = []
        current = Section(level=1, heading="", text="")
        body: list[str] = []
        chapter_no = 0

        def flush() -> None:
            text_out = "\n".join(body).strip()
            if current.heading or text_out:
                sections.append(Section(level=current.level,
                                        heading=current.heading, text=text_out,
                                        page=current.page))
            body.clear()

        def flush_table(table_elem: Any) -> None:
            grid: list[list[str]] = []
            for tr in table_elem.iter():
                if _localname(tr.tag) != "tr":
                    continue
                cells = [_xhtml_text(c) for c in tr
                         if _localname(c.tag) in ("th", "td")]
                if any(cells):
                    grid.append(cells)
            if not grid:
                return
            width = max(len(row) for row in grid)
            grid = [row + [""] * (width - len(row)) for row in grid]
            doc.tables.append(Table(name=f"Table {len(doc.tables) + 1}",
                                    headers=grid[0], rows=grid[1:],
                                    page=chapter_no))

        for href in spine_hrefs:
            chapter_no += 1
            try:
                chapter = ET.fromstring(archive.read(href))
            except ET.ParseError as exc:
                raise DocumentError(
                    f"malformed chapter {href} in .epub: {exc}") from exc
            bodies = [e for e in chapter.iter()
                      if _localname(e.tag) == "body"]
            scope = bodies[0] if bodies else chapter
            for elem in scope.iter():
                tag = _localname(elem.tag)
                if tag in _EPUB_HEADINGS:
                    flush()
                    current = Section(level=_EPUB_HEADINGS[tag],
                                      heading=_xhtml_text(elem), text="",
                                      page=chapter_no)
                elif tag == "p":
                    text = _xhtml_text(elem)
                    if text:
                        body.append(text)
                elif tag == "table":
                    flush_table(elem)
    flush()
    if not sections and not doc.tables:
        raise DocumentError(".epub contains no readable chapters")
    doc.sections = sections
    doc.metadata["chapters"] = len(spine_hrefs)
    if doc.tables:
        doc.metadata["tables"] = len(doc.tables)
    return doc


# ── ODT / ODS (stdlib-only zip + XML) ───────────────────────────────────────

_OFFICE_NS = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
_TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_TABLE_NS = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_DC_NS_URI = "http://purl.org/dc/elements/1.1/"

_ODT_OFFICE = f"{{{_OFFICE_NS}}}"
_ODT_TEXT = f"{{{_TEXT_NS}}}"
_ODT_TABLE = f"{{{_TABLE_NS}}}"
_ODT_DC = f"{{{_DC_NS_URI}}}"


def _odt_cell_text(cell: Any) -> str:
    return " ".join(cell.itertext()).strip()


def _odt_table_grid(table_elem: Any) -> list[list[str]]:
    grid: list[list[str]] = []
    for row in table_elem.iter(f"{_ODT_TABLE}table-row"):
        cells = [_odt_cell_text(c) for c in row
                 if _localname(c.tag) in ("table-cell", "covered-table-cell")]
        if any(cells):
            grid.append(cells)
    if not grid:
        return []
    width = max(len(row) for row in grid)
    return [row + [""] * (width - len(row)) for row in grid]


class _OdtFlow:
    """Accumulates ODT flow content (headings/paragraphs/tables) in order."""

    def __init__(self, doc: Document) -> None:
        self._doc = doc
        self.sections: list[Section] = []
        self._current = Section(level=1, heading="", text="")
        self._body: list[str] = []

    def flush(self) -> None:
        text_out = "\n".join(self._body).strip()
        if self._current.heading or text_out:
            self.sections.append(Section(level=self._current.level,
                                         heading=self._current.heading,
                                         text=text_out))
        self._body.clear()

    def heading(self, level: int, text: str) -> None:
        if not text:
            return
        self.flush()
        self._current = Section(level=min(max(level, 1), 6),
                                heading=text, text="")

    def para(self, text: str) -> None:
        if text:
            self._body.append(text)

    def table(self, grid: list[list[str]], name: str) -> None:
        if not grid:
            return
        self._doc.tables.append(
            Table(name=name, headers=grid[0], rows=grid[1:]))
        self._body.append(f"[{name}: {len(grid) - 1} data rows]")

    def walk(self, elem: Any) -> None:
        for child in elem:
            tag = _localname(child.tag)
            if tag == "h":
                try:
                    level = int(child.get(f"{_ODT_TEXT}outline-level", "1"))
                except (TypeError, ValueError):
                    level = 1
                self.heading(level, _odt_cell_text(child))
            elif tag == "p":
                self.para(_odt_cell_text(child))
            elif tag == "table":
                name = child.get(f"{_ODT_TABLE}name", "") or \
                    f"Table {len(self._doc.tables) + 1}"
                grid = _odt_table_grid(child)
                self.table(grid, name)
            elif tag in ("list", "list-item", "section", "table-of-content",
                         "index-body", "alphabetical-index", "toc"):
                self.walk(child)
            # Other elements (draw:frame, text:soft-page-break, …) carry no
            # flow text and are skipped on purpose.


def _parse_odt(data: bytes, doc: Document) -> Document:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"invalid .odt file: {exc}") from exc
    with archive:
        try:
            root = ET.fromstring(archive.read("content.xml"))
        except (KeyError, ET.ParseError) as exc:
            raise DocumentError(f".odt has no readable content.xml: {exc}") from exc
        try:
            meta_root = ET.fromstring(archive.read("meta.xml"))
            title_node = meta_root.find(f".//{_ODT_DC}title")
            if (title_node is not None and title_node.text
                    and title_node.text.strip()):
                doc.title = title_node.text.strip()
            creator_node = meta_root.find(f".//{_ODT_DC}creator")
            if (creator_node is not None and creator_node.text
                    and creator_node.text.strip()):
                doc.author = creator_node.text.strip()
            lang_node = meta_root.find(f".//{_ODT_DC}language")
            if (lang_node is not None and lang_node.text
                    and lang_node.text.strip()):
                doc.metadata["language"] = lang_node.text.strip()
        except (KeyError, ET.ParseError):
            # meta.xml is optional metadata; a missing/corrupt one must
            # never fail the document parse.
            _log.debug("odt meta.xml unreadable; continuing without it")
        body = root.find(f"{_ODT_OFFICE}body")
        if body is None:
            raise DocumentError(".odt content.xml has no office:body")
        text_root = body.find(f"{_ODT_OFFICE}text")
        if text_root is None:
            raise DocumentError(".odt has no office:text content")
        flow = _OdtFlow(doc)
        flow.walk(text_root)
        flow.flush()
    if not flow.sections and not doc.tables:
        raise DocumentError(".odt contains no paragraphs or tables")
    doc.sections = flow.sections
    if doc.tables:
        doc.metadata["tables"] = len(doc.tables)
    return doc


def _parse_ods(data: bytes, doc: Document) -> Document:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"invalid .ods file: {exc}") from exc
    with archive:
        try:
            root = ET.fromstring(archive.read("content.xml"))
        except (KeyError, ET.ParseError) as exc:
            raise DocumentError(f".ods has no readable content.xml: {exc}") from exc
        body = root.find(f"{_ODT_OFFICE}body")
        sheet_root = (body.find(f"{_ODT_OFFICE}spreadsheet")
                      if body is not None else None)
        if sheet_root is None:
            raise DocumentError(".ods has no office:spreadsheet content")
        for table_elem in sheet_root.iter(f"{_ODT_TABLE}table"):
            name = table_elem.get(f"{_ODT_TABLE}name", "") or \
                f"Sheet {len(doc.tables) + 1}"
            grid = _odt_table_grid(table_elem)
            if not grid:
                continue
            doc.tables.append(Table(name=name, headers=grid[0], rows=grid[1:]))
            doc.sections.append(Section(
                level=1, heading=name,
                text=(f"Worksheet {name!r}: {len(grid) - 1} data rows × "
                      f"{len(grid[0])} columns.")))
    if not doc.tables:
        raise DocumentError(".ods spreadsheet has no non-empty tables")
    doc.metadata["sheets"] = [t.name for t in doc.tables]
    return doc


_PARSERS: dict[str, Callable[..., Document]] = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "xls": _parse_xls,
    "pptx": _parse_pptx,
    "html": _parse_html,
    "markdown": _parse_markdown,
    "txt": _parse_txt,
    "rtf": _parse_rtf,
    "epub": _parse_epub,
    "odt": _parse_odt,
    "ods": _parse_ods,
    "json": _parse_json,
    "xml": _parse_xml,
}


def parse_bytes(data: bytes, *, filename: str = "", mime: str = "") -> Document:
    """Parse raw ``data`` into a Document.

    The format is detected from ``filename``'s extension first, then magic
    bytes, then ``mime``.  Raises :class:`DocumentError` on unknown or
    unparseable input — never returns an empty Document.
    """
    fmt = _sniff(data, filename, mime)
    doc = new_document(format=fmt, title=_stem(filename),
                       source=filename or mime or "")
    if fmt in ("csv", "tsv"):
        doc = _parse_delimited(data, doc, "\t" if fmt == "tsv" else ",", fmt)
    else:
        doc = _PARSERS[fmt](data, doc)
    _emit("document.parsed", {
        "doc_id": doc.id,
        "format": doc.format,
        "title": doc.title,
        "sections": len(doc.sections),
        "tables": len(doc.tables),
        "source": doc.source,
    })
    return doc


def parse_path(path: str | Path) -> Document:
    """Parse the file at ``path``.  Raises DocumentError when missing or bad."""
    file_path = Path(path)
    if not file_path.is_file():
        raise DocumentError(f"no such file: {file_path}")
    try:
        data = file_path.read_bytes()
    except OSError as exc:
        raise DocumentError(f"cannot read {file_path}: {exc}") from exc
    doc = parse_bytes(data, filename=file_path.name)
    doc.source = str(file_path)
    return doc
