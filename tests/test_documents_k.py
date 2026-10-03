"""Wave K document engine: parsing every format, converters, index.

Builds real fixtures (docx via python-docx, xlsx via openpyxl, pptx via
raw zipfile XML, pdf via core.pdf.render_pdf) and exercises the whole
nomorals.documents surface: parse_bytes/parse_path, every converter,
the inverted index, and fail-fast behaviour.
"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

try:
    from docx import Document as DocxDocument
    _HAS_DOCX = True
except ImportError:  # python-docx is an optional test dependency
    DocxDocument = None
    _HAS_DOCX = False

try:
    from openpyxl import Workbook
    _HAS_OPENPYXL = True
except ImportError:  # openpyxl is an optional test dependency
    Workbook = None
    _HAS_OPENPYXL = False

requires_docx = unittest.skipUnless(
    _HAS_DOCX, "python-docx not installed")
requires_openpyxl = unittest.skipUnless(
    _HAS_OPENPYXL, "openpyxl not installed")

from nomorals.core.pdf import render_pdf
from nomorals.documents import (
    Document,
    DocumentError,
    DocumentIndex,
    Section,
    Table,
    full_text,
    parse_bytes,
    parse_path,
    to_csv,
    to_html,
    to_markdown,
    to_pdf,
    to_text,
)

_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
_PRES = "http://schemas.openxmlformats.org/presentationml/2006/main"


def make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.core_properties.title = "Quarterly Report"
    doc.core_properties.author = "Devon"
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This report covers the <b>quarterly</b> results & outlook.")
    doc.add_heading("Details", level=2)
    doc.add_paragraph("Revenue grew in every segment.")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Segment"
    table.rows[0].cells[1].text = "Revenue"
    table.rows[1].cells[0].text = "North"
    table.rows[1].cells[1].text = "120"
    table.rows[2].cells[0].text = "South"
    table.rows[2].cells[1].text = "95"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Planets"
    ws1.append(["Name", "Moons"])
    ws1.append(["Earth", 1])
    ws1.append(["Mars", 2])
    ws2 = wb.create_sheet("Rovers")
    ws2.append(["Rover", "Year"])
    ws2.append(["Curiosity", 2012])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _slide_xml(body_inner: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<p:sld xmlns:p="{_PRES}" xmlns:a="{_DRAWING}">'
        "<p:cSld><p:spTree>"
        f"<p:sp><p:txBody><a:bodyPr/>{body_inner}</p:txBody></p:sp>"
        "</p:spTree></p:cSld></p:sld>"
    )


def _para(text: str) -> str:
    return f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p>"


def make_pptx_bytes() -> bytes:
    slide1 = _slide_xml(_para("First slide title") + _para("Opening bullet"))
    table_xml = (
        f'<a:tbl><a:tblPr/><a:tblGrid/>'
        f'<a:tr><a:tc><a:txBody>{_para("Alpha")}</a:txBody></a:tc>'
        f'<a:tc><a:txBody>{_para("Beta")}</a:txBody></a:tc></a:tr>'
        f'<a:tr><a:tc><a:txBody>{_para("1")}</a:txBody></a:tc>'
        f'<a:tc><a:txBody>{_para("2")}</a:txBody></a:tc></a:tr>'
        "</a:tbl>"
    )
    slide2 = _slide_xml(_para("Second slide") + table_xml)
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/presentation.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.'
        'presentation.main+xml"/>'
        '<Override PartName="/ppt/slides/slide1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.'
        'slide+xml"/>'
        "</Types>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("ppt/slides/slide1.xml", slide1)
        archive.writestr("ppt/slides/slide2.xml", slide2)
    return buffer.getvalue()


def make_pdf_bytes() -> bytes:
    lines = ["# Big Report", ""]
    for i in range(40):
        lines.append(f"## Section {i + 1}")
        lines.append("")
        lines.append(f"Body text for section {i + 1} with enough words to wrap.")
        lines.append("")
    return render_pdf("\n".join(lines), title="Big Report", headings=True)


class TestModel(unittest.TestCase):
    def test_dict_round_trip(self) -> None:
        doc = Document(
            id="doc1", format="txt", title="T", author="A", source="s",
            created_at=1234.5,
            sections=[Section(level=2, heading="H", text="body")],
            tables=[Table(name="T1", headers=["a"], rows=[["1"], ["2"]])],
            metadata={"k": "v"},
        )
        clone = Document.from_dict(doc.to_dict())
        self.assertEqual(clone.to_dict(), doc.to_dict())

    def test_full_text_covers_everything(self) -> None:
        doc = Document(
            title="Title",
            sections=[Section(level=1, heading="Head", text="prose here")],
            tables=[Table(name="Grid", headers=["c1"], rows=[["v1"]])],
        )
        text = full_text(doc)
        for needle in ("Title", "Head", "prose here", "Grid", "c1", "v1"):
            self.assertIn(needle, text)

    def test_from_dict_rejects_garbage(self) -> None:
        with self.assertRaises(ValueError):
            Document.from_dict({"sections": ["nope"]})


class TestParsers(unittest.TestCase):
    @requires_docx
    def test_docx_headings_and_table(self) -> None:
        doc = parse_bytes(make_docx_bytes(), filename="report.docx")
        self.assertEqual(doc.format, "docx")
        self.assertEqual(doc.title, "Quarterly Report")
        self.assertEqual(doc.author, "Devon")
        headings = [s.heading for s in doc.sections]
        self.assertIn("Introduction", headings)
        self.assertIn("Details", headings)
        intro = next(s for s in doc.sections if s.heading == "Introduction")
        self.assertEqual(intro.level, 1)
        self.assertIn("quarterly", intro.text)
        details = next(s for s in doc.sections if s.heading == "Details")
        self.assertEqual(details.level, 2)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(table.headers, ["Segment", "Revenue"])
        self.assertEqual(table.rows, [["North", "120"], ["South", "95"]])

    @requires_openpyxl
    def test_xlsx_two_sheets(self) -> None:
        doc = parse_bytes(make_xlsx_bytes(), filename="data.xlsx")
        self.assertEqual(doc.format, "xlsx")
        self.assertEqual([t.name for t in doc.tables], ["Planets", "Rovers"])
        planets = doc.tables[0]
        self.assertEqual(planets.headers, ["Name", "Moons"])
        self.assertEqual(planets.rows, [["Earth", "1"], ["Mars", "2"]])
        rovers = doc.tables[1]
        self.assertEqual(rovers.headers, ["Rover", "Year"])
        self.assertEqual([s.heading for s in doc.sections], ["Planets", "Rovers"])

    def test_pptx_two_slides(self) -> None:
        doc = parse_bytes(make_pptx_bytes(), filename="deck.pptx")
        self.assertEqual(doc.format, "pptx")
        self.assertEqual(len(doc.sections), 2)
        self.assertEqual(doc.sections[0].heading, "Slide 1")
        self.assertIn("First slide title", doc.sections[0].text)
        self.assertIn("Opening bullet", doc.sections[0].text)
        self.assertEqual(doc.sections[1].heading, "Slide 2")
        self.assertIn("Second slide", doc.sections[1].text)
        self.assertIn("Alpha | Beta", doc.sections[1].text)
        self.assertIn("1 | 2", doc.sections[1].text)

    def test_pptx_sniffed_by_magic_bytes(self) -> None:
        doc = parse_bytes(make_pptx_bytes())
        self.assertEqual(doc.format, "pptx")
        self.assertEqual(len(doc.sections), 2)

    def test_pdf_pages_become_sections(self) -> None:
        doc = parse_bytes(make_pdf_bytes(), filename="big.pdf")
        self.assertEqual(doc.format, "pdf")
        self.assertGreaterEqual(len(doc.sections), 2)
        self.assertEqual(doc.sections[0].heading, "Page 1")
        self.assertEqual(doc.sections[-1].heading,
                         f"Page {len(doc.sections)}")
        self.assertEqual(doc.metadata["pages"], len(doc.sections))
        text = full_text(doc)
        self.assertIn("Section 1", text)
        self.assertIn("Section 40", text)
        self.assertIn("Body text for section 20", text)

    def test_pdf_sniffed_by_magic_bytes(self) -> None:
        doc = parse_bytes(make_pdf_bytes())
        self.assertEqual(doc.format, "pdf")

    def test_markdown_sections_and_code_blocks(self) -> None:
        md = ("# Top\n\nintro text\n\n```\n# not a heading\n```\n\n"
              "## Second\n\nmore text\n")
        doc = parse_bytes(md.encode(), filename="notes.md")
        self.assertEqual(doc.format, "markdown")
        self.assertEqual([s.heading for s in doc.sections], ["Top", "Second"])
        self.assertEqual(doc.sections[0].level, 1)
        self.assertEqual(doc.sections[1].level, 2)
        self.assertIn("# not a heading", doc.sections[0].text)

    def test_markdown_by_mime(self) -> None:
        doc = parse_bytes(b"# Hi\n\nbody\n", mime="text/markdown")
        self.assertEqual(doc.format, "markdown")
        self.assertEqual(doc.sections[0].heading, "Hi")

    def test_html_headings_and_table(self) -> None:
        html = (
            "<html><head><title>Test Page</title></head><body>"
            "<h1>Main</h1><p>Opening paragraph.</p>"
            "<h2>Sub</h2><p>Detail &amp; more.</p>"
            "<table><tr><th>Item</th><th>Qty</th></tr>"
            "<tr><td>Apples</td><td>3</td></tr></table>"
            "</body></html>"
        )
        doc = parse_bytes(html.encode(), filename="page.html")
        self.assertEqual(doc.format, "html")
        self.assertEqual(doc.title, "Test Page")
        headings = [s.heading for s in doc.sections]
        self.assertIn("Main", headings)
        self.assertIn("Sub", headings)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].headers, ["Item", "Qty"])
        self.assertEqual(doc.tables[0].rows, [["Apples", "3"]])

    def test_csv_and_tsv(self) -> None:
        doc = parse_bytes(b"name,age\nAlice,30\nBob,25\n", filename="people.csv")
        self.assertEqual(doc.format, "csv")
        self.assertEqual(doc.tables[0].headers, ["name", "age"])
        self.assertEqual(doc.tables[0].rows, [["Alice", "30"], ["Bob", "25"]])
        tsv = parse_bytes(b"k\tv\na\t1\n", filename="data.tsv")
        self.assertEqual(tsv.format, "tsv")
        self.assertEqual(tsv.tables[0].headers, ["k", "v"])

    def test_txt_single_section(self) -> None:
        doc = parse_bytes(b"just some plain text\nsecond line", filename="n.txt")
        self.assertEqual(doc.format, "txt")
        self.assertEqual(len(doc.sections), 1)
        self.assertIn("just some plain text", doc.sections[0].text)

    def test_unknown_bytes_raise(self) -> None:
        with self.assertRaises(DocumentError):
            parse_bytes(b"\x00\x01\x02\x03\xff\xfe\x00garbage")
        with self.assertRaises(DocumentError):
            parse_bytes(b"")
        with self.assertRaises(DocumentError):
            parse_bytes(b"hello", filename="file.xyz")

    def test_corrupt_docx_raises(self) -> None:
        with self.assertRaises(DocumentError):
            parse_bytes(b"PK\x03\x04not a real zip", filename="bad.docx")

    @requires_docx
    def test_parse_path(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.docx"
            path.write_bytes(make_docx_bytes())
            doc = parse_path(path)
            self.assertEqual(doc.format, "docx")
            self.assertEqual(doc.source, str(path))
            self.assertIn("Introduction", [s.heading for s in doc.sections])

    def test_parse_path_missing_raises(self) -> None:
        with self.assertRaises(DocumentError):
            parse_path("/nonexistent/dir/file.pdf")


@requires_docx
class TestConverters(unittest.TestCase):
    def setUp(self) -> None:
        self.doc = parse_bytes(make_docx_bytes(), filename="report.docx")

    def test_to_markdown_nonempty_and_headings(self) -> None:
        md = to_markdown(self.doc)
        self.assertIn("# Introduction", md)
        self.assertIn("## Details", md)
        self.assertIn("| Segment | Revenue |", md)

    def test_to_text_nonempty(self) -> None:
        text = to_text(self.doc)
        self.assertIn("Introduction", text)
        self.assertIn("North", text)

    def test_to_html_escapes(self) -> None:
        html = to_html(self.doc)
        self.assertIn("<h1>Introduction</h1>", html)
        self.assertIn("<table>", html)
        self.assertIn("&lt;b&gt;quarterly&lt;/b&gt;", html)
        self.assertIn("&amp;", html)
        self.assertTrue(html.startswith("<!DOCTYPE html>"))

    def test_to_pdf_bytes(self) -> None:
        data = to_pdf(self.doc)
        self.assertIsInstance(data, bytes)
        self.assertTrue(data.startswith(b"%PDF"))

    def test_to_csv(self) -> None:
        csv_text = to_csv(self.doc, table=0)
        self.assertIn("Segment,Revenue", csv_text)
        self.assertIn("North,120", csv_text)

    def test_to_csv_bad_index_raises(self) -> None:
        with self.assertRaises(DocumentError):
            to_csv(self.doc, table=5)

    def test_to_csv_no_tables_raises(self) -> None:
        doc = parse_bytes(b"plain", filename="a.txt")
        with self.assertRaises(DocumentError):
            to_csv(doc)

    def test_markdown_round_trip_preserves_headings(self) -> None:
        md = to_markdown(self.doc)
        again = parse_bytes(md.encode(), filename="x.md")
        headings = {s.heading for s in again.sections}
        self.assertIn("Introduction", headings)
        self.assertIn("Details", headings)


class TestIndex(unittest.TestCase):
    def _docs(self) -> tuple[Document, Document]:
        quantum = parse_bytes(
            b"# Quantum\n\nquantum bananas quantum physics. quantum leaps.",
            filename="q.md",
        )
        carpentry = parse_bytes(
            b"# Carpentry\n\nrustic carpentry joints and timber framing.",
            filename="c.md",
        )
        return quantum, carpentry

    def test_search_ranks_correct_doc_first(self) -> None:
        high = parse_bytes(b"# High\n\napple apple apple banana", filename="h.md")
        low = parse_bytes(b"# Low\n\napple banana", filename="l.md")
        index = DocumentIndex()
        index.add(low)
        index.add(high)
        hits = index.search("apple")
        self.assertEqual([h["doc_id"] for h in hits], [high.id, low.id])
        self.assertGreater(hits[0]["score"], hits[1]["score"])
        self.assertIn("apple", hits[0]["snippet"])
        self.assertIn("apple", hits[1]["snippet"])

    def test_search_result_shape(self) -> None:
        (doc,) = self._docs()[:1]
        index = DocumentIndex()
        index.add(doc)
        (hit,) = index.search("quantum", limit=5)
        self.assertEqual(set(hit), {"doc_id", "title", "score", "snippet"})
        self.assertEqual(hit["title"], "Quantum")
        self.assertLessEqual(len(hit["snippet"]), 130)

    def test_search_empty_query_raises(self) -> None:
        index = DocumentIndex()
        with self.assertRaises(DocumentError):
            index.search("   ")

    def test_remove(self) -> None:
        quantum, _ = self._docs()
        index = DocumentIndex()
        index.add(quantum)
        self.assertTrue(index.remove(quantum.id))
        self.assertFalse(index.remove(quantum.id))
        self.assertEqual(index.search("quantum"), [])

    def test_save_load_round_trip(self) -> None:
        # NOTE (BM25 migration): save() now writes a version-2 SQLite file
        # (the FTS5 index itself is persisted) instead of version-1 JSON.
        # The format change is intentional; load() still migrates legacy v1
        # JSON — see test_load_legacy_json_v1_migrates.
        quantum, carpentry = self._docs()
        index = DocumentIndex()
        index.add(quantum)
        index.add(carpentry)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "index.db"
            index.save(path)
            self.assertEqual(path.read_bytes()[:16], b"SQLite format 3\x00")
            loaded = DocumentIndex.load(path)
        self.assertEqual(len(loaded), 2)
        hits = loaded.search("carpentry")
        self.assertEqual(hits[0]["doc_id"], carpentry.id)
        self.assertIn("carpentry", hits[0]["snippet"].lower())

    def test_load_legacy_json_v1_migrates(self) -> None:
        # A version-1 JSON index (written before the BM25 migration) loads
        # and searches on the new backend with no manual migration step.
        quantum, _ = self._docs()
        payload = {
            "version": 1,
            "docs": [{
                "id": quantum.id,
                "title": "Quantum",
                "text": full_text(quantum),
            }],
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "index.json"
            path.write_text(json.dumps(payload), encoding="utf-8")
            loaded = DocumentIndex.load(path)
        (hit,) = loaded.search("quantum")
        self.assertEqual(hit["doc_id"], quantum.id)
        self.assertIsInstance(hit["score"], float)

    def test_load_foreign_sqlite_raises(self) -> None:
        # A SQLite file that is not a document index must fail fast with
        # DocumentError, not leak storage-layer errors or return empty.
        import sqlite3 as _sqlite3

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "other.db"
            conn = _sqlite3.connect(str(path))
            conn.execute("CREATE TABLE stuff (a TEXT)")
            conn.commit()
            conn.close()
            with self.assertRaises(DocumentError):
                DocumentIndex.load(path)

    def test_load_missing_file_raises(self) -> None:
        with self.assertRaises(DocumentError):
            DocumentIndex.load("/nonexistent/index.json")


if __name__ == "__main__":
    unittest.main()
